"""Extraction where the model decides the shape, and code proves the figures.

The parsers in engine/ only read layouts somebody wrote a recipe for. Hand
them an unfamiliar report and the best they manage is a generic table dump --
which is not what was asked for, and is why this felt like code with a chat
box on top rather than an assistant.

Here the model does the work. It is given the document's text and the
requirement in the user's own words, and it returns the rows. No layout needs
to be known in advance, and the requirement can be anything: one row per
tenant, a summary by month, three named fields, whatever was asked.

The obvious objection is that a language model will invent a number. So every
value it returns is checked back against the source text before anyone sees
it. A figure that appears in the document is marked verified. One that does
not is NOT dropped and NOT silently kept -- it is written to the sheet in a
coloured cell saying it could not be found in the source. The model supplies
the shape and the judgement; the document remains the authority on the facts.

That division is the whole idea. Asking a model to be accurate is hoping.
Checking its output against the page is knowing.
"""
from __future__ import annotations

import json
import os
import re
import time

import requests

OLLAMA = os.environ.get("OLLAMA_URL", "http://127.0.0.1:11434")
MODEL = os.environ.get("FINAI_MODEL", "gpt-oss-64k:latest")

_here = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(_here, "outputs")

SCHEMA = {
    "type": "object",
    "properties": {
        "columns": {"type": "array", "items": {"type": "string"}},
        "rows": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "values": {"type": "array", "items": {"type": "string"}},
                "note": {"type": "string"},
            },
            "required": ["values"],
        }},
        "uncertain": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["columns", "rows"],
}

PROMPT = """You are extracting data from a document exactly as the user asked.

THE USER'S REQUIREMENT:
{requirement}

THE DOCUMENT:
{document}

Rules:
- Give the columns the user asked for. If they did not name columns, choose \
the ones that answer their requirement.
- One entry in "rows" per record. "values" must line up with "columns" in \
order and have the same length.
- Copy values EXACTLY as printed in the document -- same digits, same \
punctuation, same order. Do not reformat, round, convert or tidy them.
- If a value is genuinely absent for a row, use an empty string. Never guess \
and never carry a value over from another row.
- If something was ambiguous, describe it briefly in "uncertain". Being \
unsure is useful information, not a failure.

Return only the JSON."""


def _numbers_in(text):
    """Every number in a blob, normalised so 1,234.50 and 1234.5 compare equal."""
    out = set()
    for raw in re.findall(r"-?\(?\$?[\d,]*\.?\d+\)?%?-?", text):
        t = raw.strip().replace("$", "").replace(",", "").replace("%", "")
        neg = t.startswith("(") and t.endswith(")")
        t = t.strip("()").rstrip("-")
        if not t or not any(c.isdigit() for c in t):
            continue
        try:
            v = float(t)
        except ValueError:
            continue
        out.add(-v if neg else v)
    return out


def _is_number(s):
    return bool(re.fullmatch(r"-?\(?\$?[\d,]*\.?\d+\)?%?-?", (s or "").strip())
                and any(c.isdigit() for c in s))


def verify(rows, columns, source_text):
    """Check every value back against the document.

    Numbers are compared by value, so the check does not fail merely because
    the model wrote 1234.5 where the page says 1,234.50. Text is compared
    loosely on its letters and digits, because a model will normalise spacing
    and case in a name without changing who it refers to.

    Returns (flags, n_checked) where a flag names a cell that could not be
    found in the source.
    """
    src_numbers = _numbers_in(source_text)
    squashed = re.sub(r"[^a-z0-9]", "", source_text.lower())

    flags, checked = [], 0
    for i, row in enumerate(rows):
        for j, value in enumerate(row.get("values", [])):
            v = (value or "").strip()
            if not v:
                continue
            checked += 1
            if _is_number(v):
                nums = _numbers_in(v)
                if nums and not (nums & src_numbers):
                    flags.append({
                        "row": i, "column": j,
                        "column_name": columns[j] if j < len(columns) else str(j),
                        "value": v,
                        "why": "This figure does not appear anywhere in the "
                               "document text. It may have been miscopied or "
                               "invented, or it may be derived -- either way a "
                               "person should confirm it against the page.",
                    })
            else:
                probe = re.sub(r"[^a-z0-9]", "", v.lower())
                if len(probe) >= 4 and probe not in squashed:
                    flags.append({
                        "row": i, "column": j,
                        "column_name": columns[j] if j < len(columns) else str(j),
                        "value": v,
                        "why": "This text does not appear in the document. It "
                               "may be a summary or a label the model chose "
                               "rather than something the page says.",
                    })
    return flags, checked


def ask_model(requirement, document, model=MODEL, timeout=900):
    body = {
        "model": model,
        "messages": [{"role": "user",
                      "content": PROMPT.format(requirement=requirement,
                                               document=document)}],
        "stream": False,
        # The schema is enforced by the server. Asking a model in prose to
        # "return JSON" is a request; this is a guarantee.
        "format": SCHEMA,
        "options": {"temperature": 0, "top_p": 1, "top_k": 1, "seed": 7,
                    "num_ctx": 32768},
    }
    r = requests.post(f"{OLLAMA}/api/chat", json=body, timeout=timeout)
    r.raise_for_status()
    content = (r.json().get("message", {}) or {}).get("content", "")
    try:
        return json.loads(content)
    except ValueError:
        start, end = content.find("{"), content.rfind("}")
        if start >= 0 and end > start:
            return json.loads(content[start:end + 1])
        raise


def extract(path, requirement, pages=None, model=MODEL, max_chars=48000):
    """Extract whatever was asked for from any document. Returns a result dict."""
    import readers

    t0 = time.time()
    doc = readers.read_document(path, pages=pages, max_chars=max_chars)
    if doc.get("error"):
        return {"error": doc["error"]}
    text = doc.get("text") or ""
    if not text.strip():
        return {"error": "There is no readable text in that document."}

    try:
        out = ask_model(requirement, text, model=model)
    except Exception as e:
        return {"error": f"The model did not return usable JSON: "
                         f"{type(e).__name__}: {e}"}

    columns = out.get("columns") or []
    rows = out.get("rows") or []
    # A row whose values do not line up with the columns cannot be written to
    # a sheet without guessing which value belongs where, so it is padded and
    # reported rather than quietly reshaped.
    ragged = 0
    for row in rows:
        vals = row.get("values") or []
        if len(vals) != len(columns):
            ragged += 1
            row["values"] = (vals + [""] * len(columns))[:len(columns)]

    flags, checked = verify(rows, columns, text)

    # Where the text came from decides what the checking is worth. Grounding
    # compares each value against the source text -- but when that text was
    # itself transcribed from an image, the check only proves the model
    # copied its own transcription faithfully. It cannot see a digit the
    # vision model misread. On a dense table that is exactly what goes wrong,
    # so the result says so rather than reporting a clean check.
    from_vision = "vision" in (doc.get("how") or "")
    return {
        "file": os.path.basename(path),
        "how_read": doc.get("how"),
        "pages_read": doc.get("pages_read"),
        "truncated": doc.get("truncated", False),
        "columns": columns,
        "rows": rows,
        "n_rows": len(rows),
        "values_checked": checked,
        "flags": flags,
        "ragged_rows": ragged,
        "uncertain": out.get("uncertain") or [],
        "source_note": doc.get("note", ""),
        "from_vision": from_vision,
        "verification_worth": (
            "WEAK -- the text was transcribed from images, so checking a value "
            "against it only proves the transcription was copied correctly, "
            "not that it was read correctly. Every figure needs eyes on the "
            "original page." if from_vision else
            "Values were checked against the document's own text layer, which "
            "is exact."),
        "seconds": round(time.time() - t0, 1),
    }


def _as_number(v):
    """'12,666.54' -> 12666.54, but 'Suite 100' stays a string."""
    if not isinstance(v, str):
        return v
    t = v.strip().replace(",", "").replace("$", "")
    if not t:
        return v
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    try:
        n = float(t)
    except ValueError:
        return v
    n = -n if neg else n
    return int(n) if n.is_integer() and abs(n) < 1e15 else n


def to_workbook(result, path=None):
    """Write the result, colouring every value that could not be verified."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    YELLOW = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
    BOLD = Font(bold=True)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    if path is None:
        stem = os.path.splitext(result.get("file", "extract"))[0][:50]
        path = os.path.join(OUTPUT_DIR, f"{stem} - as asked.xlsx")

    columns = result["columns"]
    by_cell = {}
    for f in result["flags"]:
        by_cell[(f["row"], f["column"])] = f["why"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Extract"
    ws.append(list(columns) + ["Review Note"])
    for c in range(1, len(columns) + 2):
        ws.cell(row=1, column=c).font = BOLD

    for i, row in enumerate(result["rows"]):
        notes = [by_cell[(i, j)] for j in range(len(columns))
                 if (i, j) in by_cell]
        if row.get("note"):
            notes.append(row["note"])
        # Numbers as numbers. The model returns every value as a string,
        # because the JSON schema says so, and writing those straight into
        # cells produces a sheet that LOOKS right and cannot be used: SUM over
        # the column returns zero, sorting is alphabetical, and a date is
        # text. Anything that parses as a number is stored as one.
        ws.append([_as_number(v) for v in row.get("values", [])]
                  + ["; ".join(notes) or None])
        for j in range(len(columns)):
            if (i, j) in by_cell:
                ws.cell(row=ws.max_row, column=j + 1).fill = YELLOW
        if notes:
            ws.cell(row=ws.max_row, column=len(columns) + 1).fill = YELLOW

    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.column_dimensions[get_column_letter(len(columns) + 1)].width = 70
    ws.freeze_panes = "A2"

    info = wb.create_sheet("How this was made")
    info.append(["Source", result.get("file", "")])
    info.append(["Text came from", result.get("how_read", "")])
    info.append(["Pages read", str(result.get("pages_read", ""))])
    info.append(["Rows", result.get("n_rows", 0)])
    info.append(["Values checked against the document", result.get("values_checked", 0)])
    info.append(["Values NOT found in the document", len(result.get("flags", []))])
    info.append(["How much the checking is worth", result.get("verification_worth", "")])
    if result.get("from_vision"):
        info.append(["WARNING", "This document had no readable text layer. Every "
                                "value here was transcribed from an image by a "
                                "local model and CANNOT be treated as verified. "
                                "Check them against the original page before use."])
    if result.get("truncated"):
        info.append(["WARNING", "The document was longer than could be read in "
                                "one pass; later pages were not seen."])
    if result.get("ragged_rows"):
        info.append(["WARNING", f"{result['ragged_rows']} row(s) did not line "
                                f"up with the columns and were padded."])
    if result.get("source_note"):
        info.append(["Note", result["source_note"]])
    for u in result.get("uncertain", []):
        info.append(["Model was unsure", u])
    info.append([])
    info.append(["Every figure above was searched for in the document's own "
                 "text. Yellow cells were not found there and need a person."])
    info.column_dimensions["A"].width = 38
    info.column_dimensions["B"].width = 90

    wb.save(path)
    return path
