"""Field-skill regressions; --live also checks the installed local models."""
import io
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import workspace_extraction as extraction

BILL = [
    "Northlake Energy", "Utility type: Electricity", "Account number: 001234567",
    "Invoice number: INV-2026-08", "Customer name: Example Tenant",
    "Service address: 12 Example Road", "Meter number: 00765",
    "Billing period start: 2026-08-01", "Billing period end: 2026-08-31",
    "Invoice date: 2026-09-02", "Due date: 2026-09-20",
    "Previous reading: 10000", "Current reading: 10420",
    "Consumption: 420", "Consumption unit: kWh", "Usage charges: 84.00",
    "Fixed charges: 16.00", "Taxes: 5.00", "Current charges: 105.00",
    "Previous balance: 20.00", "Payments: 20.00", "Total due: 105.00", "Currency: USD",
]
VALUES = ["Northlake Energy"] + [line.split(": ", 1)[1] for line in BILL[1:]]


def bill_pdf():
    stream = "\n".join(f"BT /F1 11 Tf 50 {750-i*24} Td ({line}) Tj ET" for i, line in enumerate(BILL)).encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    data = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(data))
        data += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    start = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{n:010} 00000 n \n".encode() for n in offsets)
    return data + f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()


class FieldSkillTests(unittest.TestCase):
    def setUp(self):
        self.cfg = extraction.get_skill("utility-bill-extraction")
        self.output = {"columns": self.cfg["columns"], "rows": [{"values": list(VALUES)}], "uncertain": []}

    def test_leading_zero_identifiers_and_source_boundaries(self):
        pages = [{"page": 2, "text": "Account number: 001234567", "method": "Native document text"}]
        self.assertEqual(extraction.source_matches("001234567", pages)[0]["page"], 2)
        self.assertEqual(extraction.source_matches("1234567", pages), [])
        self.assertEqual(extraction.source_matches("123", pages), [])

    def test_columns_cannot_shift_silently(self):
        self.output["rows"][0]["values"].pop()
        with self.assertRaises(extraction.ExtractionError):
            extraction.validate_result(self.output, self.cfg)

    def test_mixed_documents_read_every_scanned_page(self):
        with patch.object(extraction.readers, "_read_pdf_text", return_value=({1: "Text page " * 30, 2: ""}, 300)), \
             patch("pdfplumber.open") as pdf, \
             patch.object(extraction.readers, "read_page_with_vision", return_value="Scanned page account 001234567") as vision:
            pdf.return_value.__enter__.return_value.pages = [object(), object()]
            pages = extraction.read_pages("mixed.pdf", self.cfg, {extraction.readers.VISION_MODEL}, lambda *_: None)
            self.assertEqual([p["page"] for p in pages], [1, 2])
            self.assertEqual(pages[1]["method"], "Image transcription")
            self.assertEqual(vision.call_args.args[1], 2)

    def test_workbook_evidence_and_missing_required_field(self):
        self.output["rows"][0]["values"][2] = ""
        pages = [{"page": 1, "text": "\n".join(BILL), "method": "Native document text"}]
        with patch.object(extraction, "available_models", return_value={"qwen2.5:14b-instruct-16k"}), \
             patch.object(extraction, "read_pages", return_value=pages), \
             patch.object(extraction.ai_extract, "ask_model", return_value=self.output):
            result = extraction.extract("unused.pdf", self.cfg["id"], "")
        self.assertTrue(any(i.get("destination") == "Extract!C2" and i["severity"] == "Blocking" for i in result["issues"]))
        self.assertTrue(any(c["check"] == "Printed bill balance" and c["ok"] for c in result["checks"]))
        result["rows"][0]["values"][0] = "=HYPERLINK(\"https://example.invalid\")"
        output = io.BytesIO()
        extraction.write_workbook(result, output)
        book = load_workbook(output)
        self.assertEqual(book["Extract"]["A2"].data_type, "s")
        self.assertEqual(book["Extract"]["G2"].value, "00765")
        self.assertIn("Evidence", book.sheetnames)
        self.assertIsNotNone(book["Extract"]["G2"].comment)
        book.close()


def live():
    folder = Path(__file__).resolve().parent / "outputs" / "workspace-synthetic-check"
    folder.mkdir(exist_ok=True)
    source = folder / "synthetic-utility.pdf"
    source.write_bytes(bill_pdf())
    def progress(stage, fraction):
        print(stage, flush=True)
    result = extraction.extract(source, "utility-bill-extraction", "", progress)
    row = dict(zip(result["columns"], result["rows"][0]["values"]))
    assert row["Account number"] == "001234567", row
    assert row["Total due"] == "105.00", row
    assert row["Provider"] == "Northlake Energy", row
    extraction.write_workbook(result, folder / "utility-result.xlsx")
    (folder / "utility-result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print("PASS live digital PDF: provider, leading-zero account, total due, workbook and evidence", flush=True)
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1200, 1700), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 32)
    for i, line in enumerate(BILL):
        draw.text((55, 70+i*62), line, fill="black", font=font)
    image_path = folder / "synthetic-utility.png"
    image.save(image_path)
    result = extraction.extract(image_path, "extract-requested-fields", "Extract exactly these columns: Provider, Account number, Total due. One row per bill.", progress)
    values = dict(zip(result["columns"], result["rows"][0]["values"]))
    assert values["Account number"] == "001234567", values
    assert values["Total due"] == "105.00", values
    assert result["pages"][0]["method"] == "Image transcription"
    extraction.write_workbook(result, folder / "image-result.xlsx")
    (folder / "image-result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print("PASS live image: automatic vision, requested columns, leading-zero account and total", flush=True)


if __name__ == "__main__":
    live() if "--live" in sys.argv else unittest.main(verbosity=2)
