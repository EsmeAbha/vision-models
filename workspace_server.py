"""Local document operations workspace. Run: .venv/Scripts/python workspace_server.py"""
from __future__ import annotations

import hashlib
import json
import secrets
import shutil
import sqlite3
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request, UploadFile, File
from fastapi.responses import FileResponse, JSONResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
import workspace_extraction as field_skills

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "outputs" / "workspace"
INTAKE = ROOT / "uploads" / "workspace"
for folder in (DATA, INTAKE):
    folder.mkdir(parents=True, exist_ok=True)
DB = DATA / "workspace.db"
TOKEN = secrets.token_urlsafe(32)
POOL = ThreadPoolExecutor(max_workers=1)
LOCK = threading.RLock()
app = FastAPI(docs_url=None, redoc_url=None)
app.mount("/assets", StaticFiles(directory=ROOT / "workspace_web"), name="assets")


def now():
    return datetime.now(timezone.utc).isoformat()


@contextmanager
def connection():
    con = sqlite3.connect(DB, timeout=20)
    try:
        with con:
            yield con
    finally:
        con.close()


with connection() as con:
    con.execute("CREATE TABLE IF NOT EXISTS records (kind TEXT, id TEXT, body TEXT, PRIMARY KEY(kind,id))")


def save(kind, key, body):
    with LOCK, connection() as con:
        con.execute("INSERT OR REPLACE INTO records VALUES (?,?,?)", (kind, key, json.dumps(body, default=str)))


def get(kind, key):
    with connection() as con:
        row = con.execute("SELECT body FROM records WHERE kind=? AND id=?", (kind, key)).fetchone()
    if not row:
        raise HTTPException(404, "This item is unavailable.")
    return json.loads(row[0])


def all_records(kind):
    with connection() as con:
        return [json.loads(r[0]) for r in con.execute("SELECT body FROM records WHERE kind=? ORDER BY rowid DESC", (kind,))]


def audit(event, reference=""):
    key = uuid.uuid4().hex
    save("audit", key, {"id": key, "time": now(), "actor": "Local owner", "event": event, "reference": reference})


def recover_jobs():
    for old in all_records("job"):
        if old["status"] in ("Queued", "Running"):
            old.update(status="Interrupted", stage="Service restarted", error="The service restarted. Retry this job to continue.")
            save("job", old["id"], old)


@app.middleware("http")
async def local_session(request: Request, call_next):
    if request.headers.get("host") not in ("127.0.0.1:7880", "localhost:7880", "testserver"):
        return JSONResponse({"detail": "Use the local workspace address."}, status_code=403)
    if request.url.path.startswith("/api"):
        if request.cookies.get("workspace_session") != TOKEN:
            return JSONResponse({"detail": "Open the local workspace to establish a session."}, status_code=401)
        origin = request.headers.get("origin")
        if origin and origin not in ("http://127.0.0.1:7880", "http://localhost:7880", "http://testserver"):
            return JSONResponse({"detail": "This request is not permitted."}, status_code=403)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Cache-Control"] = "no-store"
    response.headers["Content-Security-Policy"] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'"
    return response


@app.get("/", response_class=HTMLResponse)
def index():
    response = HTMLResponse((ROOT / "workspace_web" / "index.html").read_text(encoding="utf-8"))
    response.set_cookie("workspace_session", TOKEN, httponly=True, samesite="strict")
    return response


def catalog():
    from engine.skills import load_all
    items = [{"id": "table-extraction", "name": "Financial table extraction", "description": "Read digital PDF tables into a traceable workbook and check available totals.", "version": "1.0.0", "owner": "Document operations", "status": "Local", "risk": "Medium", "input": "Digital PDF", "output": "XLSX + reports", "validation": "Runtime checks", "dependencies": ["Native document reader", "Workbook writer"]},
             {"id": "workbook-qc", "name": "Workbook integrity review", "description": "Inspect workbook structure, cell errors and external links. Source-to-cell reconciliation requires an approved mapping.", "version": "1.0.0", "owner": "Quality operations", "status": "Local", "risk": "Low", "input": "PDF + XLSX", "output": "QC + exception report", "validation": "Structural checks only", "dependencies": ["Workbook reader"]}]
    for cfg in load_all():
        items.append({"id": cfg["slug"], "name": cfg.get("title", cfg["slug"]), "description": cfg.get("summary", ""), "version": "1 (repository)", "owner": "Repository", "status": "Local", "risk": "Medium", "input": "Matching digital PDF", "output": "XLSX + reports", "validation": "Reconciled on each run", "dependencies": ["Native document reader", "Workbook writer"]})
    return field_skills.skills() + items


@app.get("/api/state")
def state():
    jobs = all_records("job")
    broad = [{k: j.get(k) for k in ("id", "created", "updated", "status", "stage", "progress", "mode", "skill", "version", "user", "file_count", "artifacts", "attempt")} for j in jobs]
    return {"jobs": broad, "skills": catalog(), "drafts": all_records("draft"), "audit": all_records("audit")[:100], "roots": [{"id": "intake", "name": "Workspace intake", "path": str(INTAKE), "scope": "Read uploaded documents; direct files only"}], "destination": "Private job folder", "identity": "Local owner", "local_only": True}


@app.post("/api/upload")
async def upload(files: list[UploadFile] = File(...)):
    if not 1 <= len(files) <= 20:
        raise HTTPException(400, "Select between 1 and 20 files.")
    received = []
    for file in files:
        name = Path((file.filename or "document").replace("\\", "/")).name[:180]
        suffix = Path(name).suffix.lower()
        if suffix not in (".pdf", ".xlsx", ".png", ".jpg", ".jpeg", ".webp"):
            raise HTTPException(400, "This workspace accepts PDF, PNG, JPG, WEBP and XLSX files.")
        key = uuid.uuid4().hex
        path = INTAKE / (key + suffix)
        size = 0
        with path.open("wb") as stream:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > 30 * 1024 * 1024:
                    stream.close()
                    path.unlink(missing_ok=True)
                    raise HTTPException(400, "Each file must be 30 MB or smaller.")
                stream.write(chunk)
        item = {"id": key, "name": name, "size": size, "suffix": suffix, "path": str(path)}
        save("file", key, item)
        received.append({k: v for k, v in item.items() if k != "path"})
    audit("Files added to private intake", str(len(received)))
    return received


class PathInput(BaseModel):
    path: str = Field(max_length=500)


@app.post("/api/intake")
def intake(body: PathInput):
    path = (INTAKE / body.path).resolve()
    if not path.is_relative_to(INTAKE.resolve()) or not path.exists():
        raise HTTPException(400, "Choose an existing file or folder inside the approved intake root.")
    paths = sorted(path.iterdir()) if path.is_dir() else [path]
    paths = [p for p in paths if p.is_file() and p.suffix.lower() in (".pdf", ".xlsx", ".png", ".jpg", ".jpeg", ".webp") and p.resolve().is_relative_to(INTAKE.resolve())]
    if not 1 <= len(paths) <= 20 or any(p.stat().st_size > 30 * 1024 * 1024 for p in paths):
        raise HTTPException(400, "Choose 1–20 supported files, each at most 30 MB.")
    results = []
    for p in paths:
        key = uuid.uuid4().hex
        target = INTAKE / (key + p.suffix.lower())
        shutil.copyfile(p, target)
        item = {"id": key, "name": p.name, "size": target.stat().st_size, "suffix": p.suffix.lower(), "path": str(target)}
        save("file", key, item)
        results.append({k: v for k, v in item.items() if k != "path"})
    return results


class PlanInput(BaseModel):
    files: list[str] = Field(min_length=1, max_length=20)
    prompt: str = Field(default="", max_length=4000)
    skill: str = "table-extraction"
    mode: str = "extract"
    notes: str = Field(default="", max_length=2000)


@app.post("/api/plan")
def plan(body: PlanInput):
    files = [get("file", fid) for fid in dict.fromkeys(body.files)]
    if body.mode not in ("extract", "qc"):
        raise HTTPException(400, "Select a supported job mode.")
    skill = next((s for s in catalog() if s["id"] == body.skill), None)
    if not skill or (body.mode == "qc") != (body.skill == "workbook-qc"):
        raise HTTPException(400, "Select a preset compatible with the job mode.")
    pdfs = [f for f in files if f["suffix"] == ".pdf"]
    books = [f for f in files if f["suffix"] == ".xlsx"]
    is_field_skill = skill.get("adapter") == "workspace_extraction"
    if is_field_skill:
        if books:
            raise HTTPException(400, "Field extraction accepts PDF and document images, not prepared workbooks.")
    elif not pdfs or any(f["suffix"] not in (".pdf", ".xlsx") for f in files) or (body.mode == "extract" and books) or (body.mode == "qc" and len(books) != 1):
        raise HTTPException(400, "Add PDF sources. Quality review also requires exactly one prepared XLSX workbook; extraction accepts only PDFs.")
    prompt = body.prompt.strip() or skill.get("default_prompt", "")
    if len(prompt) < 5:
        raise HTTPException(400, "Describe the fields or records you want to extract.")
    key = uuid.uuid4().hex[:12]
    stages = ["Receive documents", "Read source tables", "Validate figures", "Build workbook", "Quality review", "Prepare results"] if body.mode == "extract" else ["Receive documents", "Inspect workbook", "Quality review", "Prepare reports"]
    item = {"id": key, "created": now(), "updated": now(), "user": "Local owner", "status": "Planned", "stage": "Awaiting confirmation", "progress": 0, "mode": body.mode, "skill": skill["name"], "skill_id": skill["id"], "version": skill["version"], "prompt": body.prompt, "notes": body.notes, "file_count": len(files), "files": [{"id": f["id"], "name": f["name"], "status": "Ready"} for f in files], "stages": stages, "artifacts": [], "findings": [], "trace": [], "attempt": 0, "scope": "The selected preset defines execution. Your task and notes are retained for review; arbitrary instructions and custom mappings are not executed.", "qc_mandatory": True}
    item["prompt"] = prompt
    if is_field_skill:
        item["stages"] = ["Receive documents", "Read document pages", "Extract requested fields", "Check extracted fields", "Build workbook", "Prepare results"]
        item["scope"] = "This skill uses the existing local document reader automatically, including image reading when needed. Your task controls the extraction. Results include source evidence and require review before use."
        item["output_fields"] = skill["columns"] or ["Fields described in your task"]
    save("job", key, item)
    audit("Plan prepared", key)
    return item


@app.get("/api/jobs/{key}")
def job(key: str):
    return get("job", key)


@app.post("/api/jobs/{key}/run")
def run(key: str):
    with LOCK:
        item = get("job", key)
        if item["status"] not in ("Planned", "Failed", "Interrupted"):
            raise HTTPException(409, "This job cannot be started in its current state.")
        item.update(status="Queued", stage="Waiting for a worker", progress=0, attempt=item["attempt"] + 1, updated=now(), findings=[], artifacts=[], trace=[], extractions=[])
        item.pop("error", None)
        save("job", key, item)
        POOL.submit(process_job, key)
    audit("Job queued", key)
    return item


def finding(item, message, file, page=None, region=None, expected=None, actual=None, destination=None, severity="Blocking"):
    item["findings"].append({"id": uuid.uuid4().hex[:10], "severity": severity, "message": message, "file": file, "page": page, "region": region, "expected": expected, "actual": actual, "destination": destination, "status": "Open", "note": ""})


def checkpoint(item, stage, progress):
    item.update(status="Running", stage=stage, progress=progress, updated=now())
    item["trace"].append({"time": now(), "stage": stage, "progress": progress})
    save("job", item["id"], item)


def artifact(item, path, kind):
    item["artifacts"].append({"id": path.name, "name": path.name, "kind": kind, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "size": path.stat().st_size})


def process_job(key):
    item = get("job", key)
    folder = DATA / key / str(item["attempt"])
    folder.mkdir(parents=True, exist_ok=True)
    checks = []
    started = time.monotonic()
    try:
        checkpoint(item, "Receive documents", 5)
        files = [get("file", f["id"]) for f in item["files"]]
        for index, f in enumerate(files):
            previous_findings = len(item["findings"])
            item["files"][index]["status"] = "Processing"
            checkpoint(item, "Inspect workbook" if item["mode"] == "qc" else "Read source tables", 10 + int(65 * index / len(files)))
            path = Path(f["path"])
            if item["mode"] == "qc":
                if f["suffix"] == ".xlsx":
                    from openpyxl import load_workbook
                    import zipfile
                    with zipfile.ZipFile(path) as z:
                        if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                            raise ValueError("Workbook expands beyond the processing limit")
                    wb = load_workbook(path, read_only=True, data_only=False, keep_links=True)
                    checks.append({"check": "Workbook opens", "ok": True, "sheets": len(wb.sheetnames)})
                    if wb._external_links:
                        finding(item, "Workbook contains external links. Confirm the referenced values.", f["name"])
                    cell_count = 0
                    for ws in wb:
                        for row in ws:
                            cell_count += len(row)
                            if cell_count > 500000:
                                raise ValueError("Workbook cell limit exceeded")
                            for c in row:
                                if c.data_type == "e":
                                    finding(item, "A workbook cell contains an error.", f["name"], actual=c.value, destination=f"{ws.title}!{c.coordinate}")
                    wb.close()
                    finding(item, "Source-to-cell reconciliation needs an approved mapping. Structural checks alone do not certify this workbook.", f["name"], expected="Approved source-to-cell mapping", actual="Not configured")
                else:
                    import pdfplumber
                    with pdfplumber.open(path) as pdf:
                        checks.append({"check": "Source document opens", "ok": True, "pages": len(pdf.pages)})
            elif item["skill_id"] in {s["id"] for s in field_skills.skills()}:
                def progress(stage, fraction):
                    checkpoint(item, stage, 10 + int(70 * (index + fraction) / len(files)))
                result = field_skills.extract(path, item["skill_id"], item["prompt"], progress)
                progress("Build workbook", 0.95)
                out = folder / f"workbook-{index + 1:02d}.xlsx"
                field_skills.write_workbook(result, out)
                artifact(item, out, "Workbook")
                structured = folder / f"extracted-fields-{index + 1:02d}.json"
                structured.write_text(json.dumps({k: v for k, v in result.items() if k != "pages"}, indent=2), encoding="utf-8")
                artifact(item, structured, "Extracted fields & evidence")
                transcript = folder / f"source-transcript-{index + 1:02d}.txt"
                transcript.write_text("\n\n".join(f"PAGE {p['page']} ({p['method']})\n{p['text']}" for p in result["pages"]), encoding="utf-8")
                artifact(item, transcript, "Source transcript")
                item.setdefault("extractions", []).append({"source": f["name"], "columns": result["columns"], "rows": result["rows"][:20], "total_rows": len(result["rows"]), "pages_read": len(result["pages"])})
                for issue in result["issues"]:
                    finding(item, issue["message"], f["name"], issue.get("page"), issue.get("region"), issue.get("expected"), issue.get("actual"), issue.get("destination"), issue["severity"])
                checks.extend(result["checks"])
            else:
                extract_file(item, f, folder, index, checks)
            item["files"][index]["status"] = "Needs review" if len(item["findings"]) > previous_findings else "Reviewed" if item["mode"] == "qc" else "Processed"
        checkpoint(item, "Quality review", 85)
        report = {"job": key, "skill": item["skill"], "version": item["version"], "checks": checks, "findings": item["findings"], "scope": "Native table extraction and available source totals" if item["mode"] == "extract" else "Structural review; source reconciliation not configured"}
        if item.get("extractions"):
            report["scope"] = "Model-assisted field extraction, literal source-text occurrence checks, required fields and applicable utility balance checks. Text occurrence is not semantic certification."
        for name, content, kind in [("quality-report.json", report, "QC report"), ("exceptions.json", item["findings"], "Exception report")]:
            path = folder / name
            path.write_text(json.dumps(content, indent=2, default=str), encoding="utf-8")
            artifact(item, path, kind)
        checkpoint(item, "Prepare results", 95)
        path = folder / "provenance.json"
        path.write_text(json.dumps({"job": key, "skill": item["skill"], "version": item["version"], "sources": [{"reference": f"Source {i+1}", "sha256": hashlib.sha256(Path(f["path"]).read_bytes()).hexdigest()} for i, f in enumerate(files)], "trace": item["trace"], "artifacts": item["artifacts"]}, indent=2), encoding="utf-8")
        artifact(item, path, "Provenance")
        item.update(status="Needs review" if item["findings"] else "Completed", stage="Review findings" if item["findings"] else "Complete", progress=100, updated=now(), runtime=round(time.monotonic() - started, 2))
        save("job", key, item)
        audit("Job finished: " + item["status"], key)
    except Exception as exc:
        safe_error = str(exc) if isinstance(exc, field_skills.ExtractionError) else "A document could not be processed within this workspace's limits. Check that it is a readable, unencrypted PDF, document image or XLSX and retry."
        item.update(status="Failed", stage="Processing stopped", updated=now(), error=safe_error)
        for f in item["files"]:
            if f["status"] == "Processing":
                f["status"] = "Failed"
        save("job", key, item)
        audit("Job failed; details withheld from analyst view", key)


def extract_file(item, f, folder, index, checks):
    from engine import geometry as G, skills as S, tabular as T, validate as V, writer as W
    import pdfplumber
    path = f["path"]
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) > 150:
            finding(item, "This document exceeds the 150-page limit. Split it into smaller files.", f["name"])
            return
    tokens = G.from_pdf(path)
    if not tokens:
        finding(item, "This document needs image-based reading, which is not connected in this workspace.", f["name"])
        return
    out = folder / f"workbook-{index + 1:02d}.xlsx"
    if item["skill_id"] != "table-extraction":
        cfg = next((s for s in S.load_all() if s["slug"] == item["skill_id"]), None)
        from docagent import find_skill
        matched, score, pages = find_skill(path, tokens=tokens)
        if not cfg or not matched or matched.get("slug") != cfg["slug"] or not pages:
            finding(item, "The source does not match the selected report preset. Select a compatible source or use financial table extraction.", f["name"])
            return
        result = S.apply(cfg, [t for t in tokens if t["page"] in pages])
        W.write_skill_workbook(str(out), cfg, result, S.series_indexes(cfg), source_name=f["name"])
        tests = result["checks"]
        for flag in result["flags"]:
            finding(item, str(flag.get("detail", "Value requires review")), f["name"], flag.get("page"), flag.get("y"), actual=flag.get("label"), destination="See workbook Issues sheet", severity="Warning")
    else:
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        tests = []
        count = 0
        for pg in sorted({t["page"] for t in tokens}):
            for table in T.read_tables(tokens, page=pg):
                rows, cols = table["rows"], table["cols"]
                if not rows or not cols:
                    continue
                count += 1
                ws = wb.create_sheet(f"Page {pg} table {count}"[:31])
                ws.append(["Source page", "Source region (y)", "Label"] + [f"Value {n+1}" for n in range(len(cols))])
                for r in rows:
                    values = [r.get("page", pg), r.get("y"), r.get("label", "")] + list(r.get("values", []))
                    for col_index, value in enumerate(values, 1):
                        cell = ws.cell(ws.max_row + 1 if col_index == 1 else ws.max_row, col_index, value)
                        if isinstance(value, str):
                            cell.data_type = "s"
                from openpyxl.styles import Font, PatternFill
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="174B43")
                ws.freeze_panes = "D2"
                ws.column_dimensions["C"].width = 44
                tests.extend(V.run_all(rows, cols))
        if not count:
            finding(item, "No structured financial table was found in this source.", f["name"])
            return
        wb.save(out)
    # Treat source text as literal workbook values, never as executable formulas.
    from openpyxl import load_workbook
    wb = load_workbook(out)
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    wb.save(out)
    wb.close()
    checks.extend(tests)
    for check in tests:
        if not check.get("ok"):
            finding(item, str(check.get("label") or check.get("note") or check.get("check") or "Figures do not reconcile"), f["name"], check.get("page"), check.get("y"), check.get("expected"), check.get("got"), "See source page / workbook row")
    if not tests:
        finding(item, "No reconcilable totals were found. This workbook requires an independent review.", f["name"])
    artifact(item, out, "Workbook")


class ReviewInput(BaseModel):
    action: str
    note: str = Field(min_length=3, max_length=1000)


@app.post("/api/jobs/{key}/findings/{fid}")
def review(key: str, fid: str, body: ReviewInput):
    with LOCK:
        item = get("job", key)
        if item["status"] not in ("Needs review", "Completed"):
            raise HTTPException(409, "Wait until processing finishes.")
        target = next((f for f in item["findings"] if f["id"] == fid), None)
        if not target or body.action not in ("Acknowledge", "Request correction"):
            raise HTTPException(400, "Choose a permitted review action.")
        target.update(status="Acknowledged" if body.action == "Acknowledge" else "Correction requested", note=body.note)
        # Acknowledgement never converts a blocking exception into a QC pass.
        save("job", key, item)
    audit("Review action: " + body.action, key)
    return item


@app.get("/api/jobs/{key}/download/{artifact_id}")
def download(key: str, artifact_id: str):
    item = get("job", key)
    entry = next((a for a in item["artifacts"] if a["id"] == artifact_id), None)
    if not entry:
        raise HTTPException(404, "Output unavailable.")
    path = DATA / key / str(item["attempt"]) / entry["id"]
    if not path.is_file():
        raise HTTPException(404, "Output unavailable.")
    audit("Output downloaded", key)
    return FileResponse(path, filename=entry["name"])


class DraftInput(BaseModel):
    name: str = Field(min_length=3, max_length=100)
    category: str
    config: dict


@app.post("/api/studio/drafts")
def draft(body: DraftInput):
    if body.category not in ("skill", "adapter", "workflow", "template", "policy"):
        raise HTTPException(400, "Unknown registry.")
    if len(json.dumps(body.config)) > 30000:
        raise HTTPException(400, "Configuration is too large.")
    key = uuid.uuid4().hex[:10]
    entry = {"id": key, "name": body.name, "category": body.category, "config": body.config, "status": "Draft", "version": "0.1.0", "owner": "Local owner", "created": now(), "validation": "Not tested"}
    save("draft", key, entry)
    audit("Studio draft saved", key)
    return entry


@app.post("/api/studio/drafts/{key}/test")
def test_draft(key: str):
    item = get("draft", key)
    cfg = item["config"]
    required = {"skill": ["input_schema", "output_schema", "validators"], "adapter": ["interface", "identity", "data_scope", "network_policy", "timeout", "health_check"], "workflow": ["steps"], "template": ["checksum", "compatibility", "environment"], "policy": ["roots", "retention_days", "egress", "max_file_mb"]}[item["category"]]
    missing = [k for k in required if k not in cfg or cfg[k] in (None, "")]
    if item["category"] == "workflow":
        expected = ["intake", "parse", "extract", "validate", "map", "qc", "export"]
        if cfg.get("steps") != expected:
            missing.append("compatible approved step order")
    item["validation"] = "Missing: " + ", ".join(missing) if missing else "Schema check passed; execution validation pending"
    item["status"] = "Draft" if missing else "Tested (schema)"
    item["test"] = {"time": now(), "passed": not missing, "scope": "Configuration schema only. No arbitrary scripts, macros or network adapters executed.", "runtime": "< 1s", "tokens": 0}
    save("draft", key, item)
    audit("Draft schema tested", key)
    return item


@app.post("/api/studio/drafts/{key}/review")
def submit_review(key: str):
    item = get("draft", key)
    if item["status"] != "Tested (schema)":
        raise HTTPException(400, "Pass the schema check before requesting review.")
    item["status"] = "In review"
    save("draft", key, item)
    audit("Release review requested", key)
    return item


if __name__ == "__main__":
    import uvicorn
    recover_jobs()
    print("Document workspace: http://127.0.0.1:7880", flush=True)
    uvicorn.run(app, host="127.0.0.1", port=7880, log_level="warning")
