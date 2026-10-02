"""A deal's workbook: proposing where fields go, and filling them over time."""
import json
import os
import shutil
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

import deal_book as db


def workbook(cells, path, sheets=None):
    book = openpyxl.Workbook()
    page = book.active
    page.title = "Deal"
    for ref, text in cells.items():
        page[ref] = text
    for name, extra in (sheets or {}).items():
        other = book.create_sheet(name)
        for ref, text in extra.items():
            other[ref] = text
    book.save(path)
    return path


def read(path, sheet="Deal"):
    book = openpyxl.load_workbook(path)
    try:
        page = book[sheet]
        return {c.coordinate: c.value for row in page.iter_rows()
                for c in row if c.value is not None}
    finally:
        book.close()


class Proposing(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="deal_")
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)
        self.path = os.path.join(self.dir, "t.xlsx")

    def propose(self, cells, fields, sheets=None):
        workbook(cells, self.path, sheets)
        return {p["field"]: p for p in db.propose_mapping(self.path, fields)}

    def test_a_label_sends_its_value_to_the_cell_beside_it(self):
        got = self.propose({"A1": "Account number", "A2": "Ending balance"},
                           ["Account number", "Ending balance"])
        self.assertEqual(got["Account number"]["cell"], "B1")
        self.assertEqual(got["Ending balance"]["cell"], "B2")
        self.assertEqual(got["Account number"]["how"], "right of the label")

    def test_a_field_the_template_never_mentions_is_absent(self):
        got = self.propose({"A1": "Account number"},
                           ["Account number", "Cap rate", "Appraiser"])
        self.assertEqual(sorted(got), ["Account number"])

    def test_the_fields_own_name_beats_a_loose_synonym(self):
        """"utility" is a synonym for Provider name, so a UTILITY section
        heading will match it. The real Provider row has to win."""
        got = self.propose({"A1": "UTILITY", "A2": "Provider"},
                           ["Provider name"])
        self.assertEqual(got["Provider name"]["label_cell"], "A2")
        self.assertEqual(got["Provider name"]["cell"], "B2")

    def test_an_exact_match_beats_one_on_another_sheet(self):
        got = self.propose({"A5": "Account number"}, ["Account number"],
                           sheets={"Tax": {"B2": "Account No:"}})
        self.assertEqual(got["Account number"]["sheet"], "Deal")
        self.assertEqual(got["Account number"]["cell"], "B5")

    def test_a_label_on_another_sheet_is_still_found(self):
        got = self.propose({"A1": "Nothing here"}, ["Account number"],
                           sheets={"Tax": {"B2": "Account No:"}})
        self.assertEqual(got["Account number"]["sheet"], "Tax")
        self.assertEqual(got["Account number"]["cell"], "C2")

    def test_a_cell_that_already_holds_something_is_left_alone(self):
        """B1 is occupied, so the value goes under the label instead."""
        got = self.propose({"A1": "Cap rate", "B1": "7.25%"}, ["Cap rate"])
        self.assertEqual(got["Cap rate"]["cell"], "A2")
        self.assertEqual(got["Cap rate"]["how"], "under the label")

    def test_two_fields_never_share_one_cell(self):
        got = self.propose({"A1": "Cap rate", "A2": "Appraiser"},
                           ["Cap rate", "Appraiser"])
        self.assertNotEqual(got["Cap rate"]["cell"], got["Appraiser"]["cell"])

    def test_one_label_serves_one_field_only(self):
        """"statement date" is a listed synonym for Billing date, so both
        fields match a template's single "Statement date" row. The closer
        match takes it and the other is dropped, rather than being handed the
        next cell along and inventing a column."""
        got = self.propose({"A1": "Statement date"},
                           ["Statement date", "Billing date"])
        self.assertEqual(sorted(got), ["Statement date"])
        self.assertEqual(got["Statement date"]["cell"], "B1")

    def test_a_trailing_colon_does_not_hide_a_label(self):
        got = self.propose({"A1": "Account Number:"}, ["Account number"])
        self.assertEqual(got["Account number"]["cell"], "B1")

    def test_a_workbook_that_will_not_open_says_which_one(self):
        bad = os.path.join(self.dir, "broken.xlsx")
        with open(bad, "wb") as fh:
            fh.write(b"not a workbook")
        with self.assertRaises(RuntimeError) as caught:
            db.propose_mapping(bad, ["Cap rate"])
        self.assertIn("broken.xlsx", str(caught.exception))


class Deals(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="deals_")
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)
        patcher = patch.object(db, "DEAL_DIR", self.dir)
        patcher.start()
        self.addCleanup(patcher.stop)

        src = os.path.join(self.dir, "src.xlsx")
        workbook({"A1": "Account number", "A2": "Ending balance",
                  "A3": "Cap rate"}, src)
        with open(src, "rb") as fh:
            self.bytes = fh.read()
        self.fields = ["Account number", "Ending balance", "Cap rate"]

    def make(self, name="Maple Avenue"):
        return db.create_deal(name, self.bytes, self.fields)

    def test_creating_a_deal_lays_out_its_folder(self):
        deal = self.make()
        self.assertEqual(deal["id"], "maple-avenue")
        for part in (db.TEMPLATE_NAME, db.WORKING_NAME, db.RECORD_NAME):
            self.assertTrue(os.path.isfile(db.deal_path(deal["id"], part)), part)
        self.assertEqual(len(deal["mapping"]), 3)
        self.assertFalse(deal["confirmed"])

    def test_a_second_deal_of_the_same_name_is_refused(self):
        self.make()
        with self.assertRaises(FileExistsError):
            self.make()

    def test_a_workbook_that_will_not_scan_leaves_no_folder(self):
        with self.assertRaises(RuntimeError):
            db.create_deal("Broken", b"not a workbook", self.fields)
        self.assertEqual(os.listdir(self.dir), ["src.xlsx"])

    def test_a_corrected_mapping_is_what_gets_stored(self):
        deal = self.make()
        fixed = [dict(row) for row in deal["mapping"]]
        fixed[0]["cell"] = "D9"
        saved = db.set_mapping(deal["id"], fixed)
        self.assertTrue(saved["confirmed"])
        self.assertEqual(saved["mapping"][0]["cell"], "D9")
        self.assertEqual(db.load_deal(deal["id"])["mapping"][0]["cell"], "D9")

    def test_a_correction_is_honoured_by_the_next_document(self):
        """The whole point of confirming a mapping: it is reused, not
        re-derived, so the same correction does not have to be made twice."""
        deal = self.make()
        fixed = [dict(row) for row in deal["mapping"]]
        moved = next(r for r in fixed if r["field"] == "Ending balance")
        moved["cell"] = "E7"
        db.set_mapping(deal["id"], fixed)

        out = db.apply_values(deal["id"], {"Ending balance": "5,940.20"},
                              source="statement.pdf")
        self.assertEqual([w["cell"] for w in out["written"]], ["E7"])
        cells = read(db.deal_path(deal["id"], db.WORKING_NAME))
        self.assertEqual(cells["E7"], "5,940.20")
        self.assertNotIn("B2", cells, "it still wrote to the proposed cell")

    def test_a_mapping_pointing_at_nonsense_is_refused(self):
        deal = self.make()
        bad = [dict(deal["mapping"][0], cell="over there")]
        with self.assertRaises(ValueError):
            db.set_mapping(deal["id"], bad)

    def test_a_field_added_by_hand_is_kept_and_used(self):
        """The scan only knows the field names the document types list, so a
        template with its own labels needs rows added by hand."""
        deal = self.make()
        extended = [dict(r) for r in deal["mapping"]]
        extended.append({"field": "Meter number", "cell": "B9",
                         "sheet": "Deal", "how": "set by hand"})
        saved = db.set_mapping(deal["id"], extended)
        self.assertIn("Meter number", [r["field"] for r in saved["mapping"]])

        out = db.apply_values(deal["id"], {"Meter number": "MTR-5520918"},
                              source="bill.pdf")
        self.assertEqual([w["cell"] for w in out["written"]], ["B9"])
        cells = read(db.deal_path(deal["id"], db.WORKING_NAME))
        self.assertEqual(cells["B9"], "MTR-5520918")

    def test_a_row_with_no_field_name_is_refused(self):
        deal = self.make()
        bad = [dict(deal["mapping"][0], field="  ")]
        with self.assertRaises(ValueError):
            db.set_mapping(deal["id"], bad)

    def test_the_same_field_twice_is_refused(self):
        deal = self.make()
        row = dict(deal["mapping"][0])
        with self.assertRaises(ValueError):
            db.set_mapping(deal["id"], [row, dict(row, cell="Z9")])

    def test_values_land_in_the_working_copy(self):
        deal = self.make()
        out = db.apply_values(deal["id"],
                              {"Account number": "4021587390",
                               "Ending balance": "5,940.20"},
                              source="statement.pdf")
        self.assertEqual(len(out["written"]), 2)
        cells = read(db.deal_path(deal["id"], db.WORKING_NAME))
        self.assertEqual(cells["B1"], "4021587390")
        self.assertEqual(cells["B2"], "5,940.20")
        # Cap rate was not in this document, and is reported rather than guessed.
        self.assertEqual([s["field"] for s in out["skipped"]], ["Cap rate"])

    def test_the_supplied_template_is_never_written_to(self):
        deal = self.make()
        db.apply_values(deal["id"], {"Account number": "X"}, source="a.pdf")
        with open(db.deal_path(deal["id"], db.TEMPLATE_NAME), "rb") as fh:
            self.assertEqual(fh.read(), self.bytes)

    def test_a_later_document_fills_what_an_earlier_one_left_empty(self):
        """The point of a deal: several documents, one workbook."""
        deal = self.make()
        db.apply_values(deal["id"], {"Account number": "4021587390"},
                        source="statement.pdf")
        out = db.apply_values(deal["id"], {"Cap rate": "7.25%"},
                              source="appraisal.pdf")
        self.assertEqual([w["field"] for w in out["written"]], ["Cap rate"])
        cells = read(db.deal_path(deal["id"], db.WORKING_NAME))
        self.assertEqual(cells["B1"], "4021587390")
        self.assertEqual(cells["B3"], "7.25%")

    def test_two_documents_disagreeing_is_reported_not_resolved(self):
        deal = self.make()
        db.apply_values(deal["id"], {"Ending balance": "5,940.20"},
                        source="statement.pdf")
        out = db.apply_values(deal["id"], {"Ending balance": "9,999.99"},
                              source="other.pdf")
        self.assertEqual(out["written"], [])
        self.assertEqual(len(out["clashed"]), 1)
        clash = out["clashed"][0]
        self.assertEqual(clash["kept"], "5,940.20")
        self.assertEqual(clash["offered"], "9,999.99")
        self.assertEqual(clash["from"], "statement.pdf")
        # The first answer stands until a person decides otherwise.
        cells = read(db.deal_path(deal["id"], db.WORKING_NAME))
        self.assertEqual(cells["B2"], "5,940.20")

    def test_the_same_value_twice_is_not_a_clash(self):
        deal = self.make()
        db.apply_values(deal["id"], {"Cap rate": "7.25%"}, source="a.pdf")
        out = db.apply_values(deal["id"], {"Cap rate": "7.25%"}, source="b.pdf")
        self.assertEqual(out["clashed"], [])
        self.assertEqual(out["written"], [])

    def test_every_document_is_recorded(self):
        deal = self.make()
        db.apply_values(deal["id"], {"Account number": "X"}, source="a.pdf")
        db.apply_values(deal["id"], {"Cap rate": "1%"}, source="b.pdf")
        history = db.load_deal(deal["id"])["history"]
        self.assertEqual([h["source"] for h in history], ["a.pdf", "b.pdf"])

    def test_listing_skips_a_folder_that_is_not_a_deal(self):
        self.make()
        os.makedirs(os.path.join(self.dir, "stray"), exist_ok=True)
        self.assertEqual([d["id"] for d in db.list_deals()], ["maple-avenue"])

    def test_a_record_that_will_not_parse_costs_only_itself(self):
        deal = self.make()
        os.makedirs(os.path.join(self.dir, "broken"), exist_ok=True)
        with open(os.path.join(self.dir, "broken", db.RECORD_NAME), "w") as fh:
            fh.write("{ not json")
        self.assertEqual([d["id"] for d in db.list_deals()], [deal["id"]])


if __name__ == "__main__":
    unittest.main(verbosity=2)
