"""Run the rating-guide locator over a zip of appraisals and write a workbook.

The locator in mhc_rating_table handles one file. This walks an unpacked
archive, calls it once per PDF, and lays the answers out three ways: an index
of what happened to every file, one long sheet that puts every appraisal's
guide side by side, and a verbatim copy of each table on its own sheet.

The long sheet is the point of the exercise -- comparing the same row across
thirty appraisals is the thing a per-file sheet cannot do -- but the verbatim
copies stay, because a reviewer who doubts a value needs to see the table as
it was printed, not only as it was reshaped.
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import mhc_rating_table as M
import save_output

NAVY = "1F3864"
HEAD = "DCE2F0"
WARN = "FCE4D6"

STATUS_NOTE = {
    "found": "the guide was located and read",
    "not_found": "no table in this file matched the guide",
    "scanned": "no text layer; OCR is required before this file can be read",
    "error": "the file could not be opened",
}


def find_pdfs(root):
    """Every PDF under root, as (absolute path, path shown to the reader).

    Sorted so a run over the same archive twice lists the files in the same
    order, which is what makes two workbooks comparable.
    """
    out = []
    for folder, _dirs, files in os.walk(root):
        for f in files:
            if f.lower().endswith(".pdf") and not f.startswith("."):
                full = os.path.join(folder, f)
                out.append((full, os.path.relpath(full, root).replace("\\", "/")))
    out.sort(key=lambda p: p[1].lower())
    return out


def run(root, progress=None):
    """Extract from every PDF under root. progress(done, total, name, result)."""
    pdfs = find_pdfs(root)
    results = []
    for i, (path, shown) in enumerate(pdfs, 1):
        res = M.extract_from_pdf(path, display=shown)
        results.append(res)
        if progress:
            progress(i, len(pdfs), shown, res)
    return results


def summarise(results):
    counts = {k: 0 for k in STATUS_NOTE}
    for r in results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    return counts


def _fit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header(ws, row, labels, fill=HEAD):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _index_sheet(wb, results):
    ws = wb.create_sheet("Index", 0)
    _header(ws, 1, ["File", "Status", "Page", "Match score", "Pages in file",
                    "Note"])
    for r in results:
        ws.append([r["file"], r["status"], r["page"], r["score"] or None,
                   r["pages"] or None, r["note"] or STATUS_NOTE.get(r["status"], "")])
        if r["status"] != "found":
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=WARN)
    _fit(ws, [52, 12, 7, 13, 13, 60])
    ws.freeze_panes = "A2"
    return ws


def _combined_sheet(wb, results):
    """One row per (appraisal, category), so the files line up for comparison."""
    ws = wb.create_sheet("Combined")
    _header(ws, 1, ["File", "Page", "Category", "In the standard guide"]
            + M.CLASS_COLUMNS)
    for r in results:
        if r["status"] != "found":
            continue
        for row in r["rows"]:
            if row["section"]:
                continue
            values = dict(zip(r["columns"], row["values"]))
            ws.append([r["file"], r["page"], row["category"],
                       "yes" if row["canonical"] else "no"]
                      + [values.get(c, "") for c in M.CLASS_COLUMNS])
    _fit(ws, [42, 6, 24, 18, 26, 26, 26, 26])
    ws.freeze_panes = "E2"
    for row in ws.iter_rows(min_row=2, min_col=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def _table_sheet(wb, res, used):
    """One appraisal's guide, laid out the way it was printed."""
    ws = wb.create_sheet(save_output._sheet_name(res["file"], used))
    title = res["title"] or "Rating guide"
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=1 + len(res["columns"]))
    c = ws.cell(row=1, column=1)
    c.font = Font(bold=True, color="FFFFFF", size=12)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")

    _header(ws, 2, ["Category"] + res["columns"])
    for row in res["rows"]:
        if row["section"]:
            ws.append([row["category"]])
            ws.merge_cells(start_row=ws.max_row, start_column=1,
                           end_row=ws.max_row, end_column=1 + len(res["columns"]))
            s = ws.cell(row=ws.max_row, column=1)
            s.font = Font(bold=True, color="FFFFFF")
            s.fill = PatternFill("solid", fgColor=NAVY)
            s.alignment = Alignment(horizontal="center")
            continue
        ws.append([row["printed"]] + list(row["values"]))
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    _fit(ws, [24] + [26] * len(res["columns"]))
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    foot = ws.max_row + 2
    ws.cell(row=foot, column=1,
            value=f"Source: {res['file']} page {res['page']} "
                  f"(match score {res['score']}, found by "
                  f"{res['strategy']} ruling)").font = Font(italic=True,
                                                            color="666666")
    return ws


def write_workbook(results, out_path):
    """Index + Combined + one sheet per appraisal. Returns out_path."""
    wb = Workbook()
    wb.remove(wb.active)
    _index_sheet(wb, results)
    _combined_sheet(wb, results)
    used = {"Index", "Combined"}
    for r in results:
        if r["status"] == "found":
            _table_sheet(wb, r, used)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path
