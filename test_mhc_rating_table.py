"""Check the rating-guide locator against PDFs built to look like the real thing.

The fixtures are generated rather than real, but they are no longer guesses.
The first few were written before any appraisal was to hand -- a document
carrying the guide, one carrying decoys, one with no text layer -- and those
proved the locator picks the right table, refuses the wrong ones and reports a
scan instead of guessing.

Then a real archive of fifty-three appraisals went through, and fifteen of
them failed. Neither of the two ways pdfplumber reads a table could manage
them: the guide is drawn without ruling the PDF layer exposes, so the ruled
reader finds nothing, and the whole-page text reader then smears the prose
above the table through the same columns. Two further details did the rest of
the damage -- cells wrap into lines that straddle their own row label, and
some files carry a text matrix a rounding error off true, which makes
pdfplumber call every character non-upright and read the table downwards, a
letter to a line.

build_unruled reproduces all of that, with the geometry measured off
cmh261413 cardinal.pdf page 43. Those checks are the ones that would catch a
regression in the character rebuild.

    .\\.venv\\Scripts\\python.exe test_mhc_rating_table.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

import mhc_rating_table as M

NAVY = colors.HexColor("#1F3864")
CLASS_COLUMNS_FIXTURE = ["Class A", "Class B", "Class C", "Unratable"]

# The guide exactly as the sample prints it, typos in the heading included.
TITLE = "Manufacture Housing Communites Rating Guide"
GUIDE = [
    ["Category", "Class A", "Class B", "Class C", "Unratable"],
    ["Density", "Low (4-7 sites/acre)", "Medium (10 sites/acre or less)",
     "High (15 sites/acre or less)", "High (15+ sites per acre)"],
    ["Age", "Built Since 1980", "Built Since 1970", "Built Prior to 1980",
     "Built Prior to 1980"],
    ["Amenities", "Resort Style", "Standard to None (if high enough quality)",
     "Few to none", "None"],
    ["Quality/Layout", "Subdivision Quality",
     "High Quality grid or curvilinear layout",
     "Medium to Low Quality/ Typical Grid Layout", "Grid Layout"],
    ["Roads", "Asphalt or Concrete", "Asphalt or Concrete",
     "Mostly Asphalt or Concrete (Some Gravel & Dirt)", "All Gravel & Dirt"],
    ["Utilities", "Public Utilities", "Usually Public Utilities",
     "Mix of Public and Private", "Mix of Public and Private"],
    ["Parking", "2 Off-Street per Site", "1 to 2 Off Street per Site",
     "Mix of Off Street and On Road", "On Road Parking Only"],
    ["Homes (Quality)", "Excellent", "Good to Excellent", "Average to Good",
     "Fair"],
    ["Homes (Age)", "Built After 1980", "Majority Built after 1980",
     "Built after 1976", "Built before 1976"],
    ["Homes (Mix)", "Mostly Multi-Section", "Single and Multisection",
     "Primarily Single Sectional", "Primarily Single Sectional"],
    ["Comparison to Star Rating", "", "", "", ""],
    ["Star Rating (Non-Woodall)", "Five Star", "Three to Four Star",
     "Two to Three Star", "One Star or Unratable"],
    ["Star Rating (Woodall)", "N/A", "N/A", "N/A", "N/A"],
]
BAND_ROWS = [11]          # "Comparison to Star Rating" spans the full width

# A table that name-drops the class columns but is not the guide. The floor has
# to reject this, or every appraisal would "find" its rent comparison instead.
DECOY = [
    ["Comparable", "Class A", "Class B", "Class C", "Unratable"],
    ["Rent", "$1,450", "$1,180", "$960", "n/a"],
    ["Occupancy", "96%", "91%", "84%", "n/a"],
    ["Cap Rate", "5.25%", "6.00%", "7.10%", "n/a"],
]

# The dangerous near-miss: an amenity comparison under the same class
# headings, borrowing three of the guide's row names. Four headers put it at
# 0.30 and three rows add 0.165, so it clears the 0.45 score floor on its own.
# Only the "half the rows or it is not the guide" gate turns it away.
NEAR_MISS = [
    ["Amenity", "Class A", "Class B", "Class C", "Unratable"],
    ["Age", "0-5 years", "5-15 years", "15-30 years", "30+ years"],
    ["Parking", "Garage", "Carport", "Surface", "Street"],
    ["Utilities", "Included", "Partly included", "Tenant paid", "Tenant paid"],
]

PROSE = ("a decade after the guide ceased publication. There is no standard "
         "rating system. Many have their own rating criteria that they "
         "reference. In this appraisal we utilize an alphabetic classification "
         "system. The highest classified buildings are Class A.")


def _grid(data, bands=()):
    """A ruled table whose cells wrap, the way Word writes one.

    The cells have to be Paragraphs, not bare strings: reportlab does not wrap
    a string, it lets it run past the column edge, and then the tail of a long
    value is clipped or lands in the next column. Real appraisals wrap -- the
    sample shows "Standard to None (if high enough quality)" on two lines --
    so a fixture built from strings would be testing a document shape that
    does not occur, and would miss whether multi-line cells survive.
    """
    body = _style("cell", 7)
    band = _style("band", 8, color=colors.white, align=1)
    grid = [[Paragraph(str(c or ""), band if r in bands else body)
             for c in row] for r, row in enumerate(data)]

    style = [
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#808080")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE2F0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for r in bands:
        style += [("SPAN", (0, r), (-1, r)),
                  ("BACKGROUND", (0, r), (-1, r), NAVY)]
    t = Table(grid, colWidths=[90, 95, 110, 110, 95], repeatRows=0)
    t.setStyle(TableStyle(style))
    return t


_STYLES = {}


def _style(name, size, color=colors.black, align=0):
    from reportlab.lib.styles import ParagraphStyle
    if name not in _STYLES:
        _STYLES[name] = ParagraphStyle(name, fontName="Helvetica",
                                       fontSize=size, leading=size + 1.5,
                                       textColor=color, alignment=align)
    return _STYLES[name]


def _titled(data, bands):
    """The guide with its merged navy heading on top, as the sample has it."""
    data = [[TITLE, "", "", "", ""]] + data
    return _grid(data, bands=[0] + [b + 1 for b in bands])


def build_with_guide(path):
    ss = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    doc.build([
        Paragraph("Appraisal of Sunset Ridge MHC", ss["Title"]),
        Paragraph(PROSE, ss["BodyText"]),
        PageBreak(),
        Paragraph("Rent Comparison", ss["Heading2"]),
        _grid(DECOY),
        PageBreak(),
        Paragraph("Market Analysis", ss["Heading1"]),
        Paragraph(PROSE, ss["BodyText"]),
        Spacer(1, 10),
        _titled(GUIDE, BAND_ROWS),
    ])


def build_without_guide(path):
    ss = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    doc.build([
        Paragraph("Appraisal of Cedar Flats Apartments", ss["Title"]),
        Paragraph(PROSE, ss["BodyText"]),
        PageBreak(),
        Paragraph("Rent Comparison", ss["Heading2"]),
        _grid(DECOY),
    ])


def build_near_miss(path):
    ss = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    doc.build([
        Paragraph("Appraisal of Willow Bend", ss["Title"]),
        Paragraph(PROSE, ss["BodyText"]),
        PageBreak(),
        Paragraph("Amenity Comparison", ss["Heading2"]),
        _grid(NEAR_MISS),
    ])


# How the real appraisals actually print the guide, which is nothing like a
# reportlab table: no ruling the PDF layer exposes, and cells that wrap into
# lines straddling their own row label -- the first line of "Medium (10
# sites/acre or less)" is printed ABOVE the word "Density", not below it. The
# geometry here is measured off cmh261413 cardinal.pdf page 43.
WRAPPED = [
    ("Density", [["Low (4-7 sites/acre)"], ["Medium (10 sites/acre or", "less)"],
                 ["High (15 sites/acre or less)"], ["High (15+ sites per acre)"]]),
    ("Age", [["Built Since 1980"], ["Built Since 1970"],
             ["Built Prior to 1980"], ["Built Prior to 1980"]]),
    ("Amenities", [["Resort Style"], ["Standard to None (if high", "enough quality)"],
                   ["Few to none"], ["None"]]),
    ("Quality/Layout", [["Subdivision Quality"],
                        ["High Quality grid or", "curvilinear layout"],
                        ["Medium to Low Quality/", "Typicall Grid Layout"],
                        ["Grid Layout"]]),
    ("Roads", [["Asphalt or Concrete"], ["Asphalt or Concrete"],
               ["Mostly Asphalt or Concrete", "(Some Gravel & Dirt)"],
               ["All Gravel & Dirt"]]),
    ("Utilities", [["Public Utilities"], ["Usually Public Utilities"],
                   ["Mix of Public and Private"], ["Mix of Public and Private"]]),
    ("Parking", [["2 Off-Street per Site"], ["1 to 2 Off Street per Site"],
                 ["Mix of Off Street and On", "Road"], ["On Road Parking Only"]]),
    ("Homes (Quality)", [["Excellent"], ["Good to Excellent"],
                         ["Average to Good"], ["Fair"]]),
    ("Homes (Age)", [["Built After 1980"], ["Majority Built after 1980"],
                     ["Built after 1976"], ["Built before 1976"]]),
    ("Homes (Mix)", [["Mostly Multi-Section"], ["Single and Multisection"],
                     ["Primarily Single Sectional"], ["Primarily Single Sectional"]]),
]
STAR_VALUES = [["Five Star"], ["Three to Four Star"], ["Two to Three Star"],
               ["One Star or Unratable"]]

CENTRES = [198.0, 306.0, 414.0, 522.0]   # the four class columns
LABEL_X = 38.2
HEADER_TOP = 197.7
PITCH = 25.66
LEADING = 12.4
PAGE_H = 792.0


def build_unruled(path, skew=False):
    """The guide as these appraisals really draw it: no ruling, wrapped cells.

    With skew=True the table is drawn through a text matrix that is very
    slightly off true, which makes pdfplumber call every character non-upright
    and read the table as vertical text, one letter to a line. Seven of the
    real files are like this and none could be read until the extractor
    stopped trusting extract_words.

    The real files are 4.9e-08 off; this uses 1e-05, because reportlab writes
    the matrix to six decimal places and anything smaller is rounded back to a
    clean zero before it reaches the file. The condition being reproduced is
    the same one either way: correct glyph positions that pdfplumber declares
    non-upright.
    """
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 9)
    c.drawString(72, PAGE_H - 70, "Appraisal of Willow Bend MHC")
    for i, chunk in enumerate(PROSE[i:i + 95] for i in range(0, 380, 95)):
        c.drawString(72, PAGE_H - 95 - i * 12, chunk)
    c.showPage()

    c.setFont("Helvetica", 9)
    c.drawString(36, PAGE_H - 60, "Market Analysis")
    for i, chunk in enumerate(PROSE[i:i + 95] for i in range(0, 380, 95)):
        c.drawString(36, PAGE_H - 90 - i * 12, chunk)
    c.drawCentredString(306, PAGE_H - 175, TITLE)

    c.saveState()
    if skew:
        c.transform(1, -1e-05, -2e-06, 1, 0, 0)
    c.setFont("Helvetica", 8)

    def put(top, x, text, centred=True):
        y = PAGE_H - top
        (c.drawCentredString if centred else c.drawString)(x, y, text)

    put(HEADER_TOP, LABEL_X, "Category", centred=False)
    for x, name in zip(CENTRES, CLASS_COLUMNS_FIXTURE):
        put(HEADER_TOP, x, name)

    for i, (label, cells) in enumerate(WRAPPED):
        top = HEADER_TOP + 20.5 + i * PITCH
        put(top, LABEL_X, label, centred=False)
        for x, lines in zip(CENTRES, cells):
            for j, line in enumerate(lines):
                put(top + (j - (len(lines) - 1) / 2) * LEADING, x, line)

    base = HEADER_TOP + 20.5 + len(WRAPPED) * PITCH
    c.setFont("Helvetica-Bold", 11)
    put(base + 6, 306, "Comparison to Star Rating")
    c.setFont("Helvetica", 8)

    # The label that wraps, with its values on the line between its two halves.
    put(base + 20.7, LABEL_X, "Star Rating (Non-", centred=False)
    put(base + 33.2, LABEL_X, "Woodall)", centred=False)
    for x, lines in zip(CENTRES, STAR_VALUES):
        put(base + 27.0, x, lines[0])

    put(base + 53.0, LABEL_X, "Star Rating (Woodall)", centred=False)
    for x in CENTRES:
        put(base + 53.0, x, "N/A")

    c.restoreState()
    c.showPage()
    c.save()


def build_scan(path):
    """A PDF with no text layer at all -- a drawn box and nothing else."""
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=letter)
    c.rect(72, 500, 400, 200, stroke=1, fill=0)
    c.rect(90, 540, 360, 30, stroke=1, fill=0)
    c.showPage()
    c.save()


FAILED = []


def check(label, got, want):
    ok = got == want
    print(f"  {'PASS' if ok else 'FAIL'}  {label}")
    if not ok:
        print(f"        wanted {want!r}")
        print(f"        got    {got!r}")
        FAILED.append(label)


def cell(res, category, column):
    """The value at one (row, column) of the extracted guide."""
    row = next((r for r in res["rows"] if r["category"] == category), None)
    if row is None:
        return f"<no row {category!r}>"
    if column not in res["columns"]:
        return f"<no column {column!r}>"
    return row["values"][res["columns"].index(column)]


def batch(tmp):
    """A zip of several appraisals, through the batch layer to a workbook."""
    import zipfile

    import mhc_batch as B
    import zips
    from openpyxl import load_workbook

    print("\na zip of several appraisals, end to end")
    src = tmp / "archive"
    (src / "2026" / "nested").mkdir(parents=True, exist_ok=True)
    build_with_guide(src / "sunset_ridge.pdf")
    build_with_guide(src / "2026" / "nested" / "oak_meadows.pdf")
    build_without_guide(src / "2026" / "cedar_flats.pdf")
    build_scan(src / "scanned_copy.pdf")
    (src / "notes.txt").write_text("not a PDF; must be ignored")

    archive = tmp / "appraisals.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        for p in src.rglob("*"):
            if p.is_file():
                zf.write(p, p.relative_to(src))

    root, _note = zips.extract(str(archive), dest_dir=str(tmp / "unpacked"))
    found = B.find_pdfs(root)
    check("the four PDFs are found, the .txt is not", len(found), 4)
    check("nested files keep their path",
          "2026/nested/oak_meadows.pdf" in [s for _p, s in found], True)

    results = B.run(root)
    counts = B.summarise(results)
    check("two files carry the guide", counts["found"], 2)
    check("one has no guide", counts["not_found"], 1)
    check("one is a scan", counts["scanned"], 1)

    out = tmp / "out" / "guides.xlsx"
    B.write_workbook(results, str(out))
    check("the workbook was written", out.exists(), True)

    wb = load_workbook(out)
    check("Index and Combined lead the workbook", wb.sheetnames[:2],
          ["Index", "Combined"])
    check("one extra sheet per appraisal that had the guide",
          len(wb.sheetnames), 4)

    ws = wb["Combined"]
    head = [c.value for c in ws[1]]
    check("Combined is keyed by file and category", head[:4],
          ["File", "Page", "Category", "In the standard guide"])
    check("and carries the four class columns", head[4:], M.CLASS_COLUMNS)

    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    check("12 guide rows for each of the 2 files", len(rows), 24)
    check("both appraisals appear", len({r[0] for r in rows}), 2)

    dens = next(r for r in rows if r[2] == "Density")
    check("a value lands in the right column", dens[4], "Low (4-7 sites/acre)")
    star = next(r for r in rows if r[2] == "Star Rating (Woodall)")
    check("the star rows made it across", star[7], "N/A")
    check("the index names every file, guide or not",
          len(list(wb["Index"].iter_rows(min_row=2))), 4)


def main():
    tmp = Path(tempfile.mkdtemp(prefix="mhc_test_"))
    good, bad, scan = tmp / "good.pdf", tmp / "bad.pdf", tmp / "scan.pdf"
    build_with_guide(good)
    build_without_guide(bad)
    build_scan(scan)
    print(f"fixtures in {tmp}\n")

    print("an appraisal that contains the guide")
    res = M.extract_from_pdf(str(good))
    check("status", res["status"], "found")
    check("found on the page the guide is on", res["page"], 3)
    check("the four class columns", res["columns"], M.CLASS_COLUMNS)
    check("heading captured as printed", res["title"], TITLE)

    labels = [r["category"] for r in res["rows"] if not r["section"]]
    for want in M.ALL_ROWS:
        check(f"row {want!r} present", want in labels, True)
    check("the star-rating band is kept as a section",
          any(r["section"] and "Star Rating" in r["category"]
              for r in res["rows"]), True)

    check("Density / Class A", cell(res, "Density", "Class A"),
          "Low (4-7 sites/acre)")
    check("Amenities / Class B", cell(res, "Amenities", "Class B"),
          "Standard to None (if high enough quality)")
    check("Roads / Class C", cell(res, "Roads", "Class C"),
          "Mostly Asphalt or Concrete (Some Gravel & Dirt)")
    check("Homes (Mix) / Unratable", cell(res, "Homes (Mix)", "Unratable"),
          "Primarily Single Sectional")
    check("Star Rating (Non-Woodall) / Class A",
          cell(res, "Star Rating (Non-Woodall)", "Class A"), "Five Star")
    check("Star Rating (Woodall) / Unratable",
          cell(res, "Star Rating (Woodall)", "Unratable"), "N/A")
    check("nothing from the decoy table leaked in",
          any("$1,450" in v for r in res["rows"] for v in r["values"]), False)

    print("\nan appraisal with a convincing decoy but no guide")
    res = M.extract_from_pdf(str(bad))
    check("status", res["status"], "not_found")
    check("no rows handed back", res["rows"], [])
    check("the decoy scored below the floor", res["score"] < M.MIN_SCORE, True)

    for skew in (False, True):
        kind = ("drawn with no ruling, cells wrapping around their labels"
                if not skew else
                "the same, through a text matrix a rounding error off true")
        print(f"\n{kind}")
        f = tmp / f"unruled_{int(skew)}.pdf"
        build_unruled(f, skew=skew)

        if skew:
            import pdfplumber
            with pdfplumber.open(f) as pdf:
                upright = {c.get("upright") for c in pdf.pages[1].chars}
            check("pdfplumber calls the characters non-upright",
                  False in upright, True)

        res = M.extract_from_pdf(str(f))
        check("status", res["status"], "found")
        check("read by rebuilding from the characters", res["strategy"], "chars")
        labels = [r["category"] for r in res["rows"] if not r["section"]]
        check("all twelve rows", labels, M.ALL_ROWS)
        check("the star band survives",
              any(r["section"] for r in res["rows"]), True)

        # The values that straddle their label are the ones that go wrong when
        # a wrapped line is folded into the row above or below it.
        check("a cell wrapped above its own label",
              cell(res, "Density", "Class B"), "Medium (10 sites/acre or less)")
        check("Amenities / Class B", cell(res, "Amenities", "Class B"),
              "Standard to None (if high enough quality)")
        check("Roads / Class C", cell(res, "Roads", "Class C"),
              "Mostly Asphalt or Concrete (Some Gravel & Dirt)")
        check("Parking / Class C", cell(res, "Parking", "Class C"),
              "Mix of Off Street and On Road")
        check("the row beneath it is not polluted",
              cell(res, "Utilities", "Class C"), "Mix of Public and Private")
        check("a label that wraps is put back together",
              cell(res, "Star Rating (Non-Woodall)", "Unratable"),
              "One Star or Unratable")
        check("Star Rating (Woodall) / Class A",
              cell(res, "Star Rating (Woodall)", "Class A"), "N/A")

    print("\na table that borrows three of the guide's rows")
    near = tmp / "near.pdf"
    build_near_miss(near)
    res = M.extract_from_pdf(str(near))
    check("status", res["status"], "not_found")
    check("it did clear the score floor on its own",
          res["score"] >= M.MIN_SCORE, True)
    check("but not the row gate", res["rows"], [])
    check("and the note says why", "of the guide's" in res["note"], True)

    print("\na scanned appraisal with no text layer")
    res = M.extract_from_pdf(str(scan))
    check("status", res["status"], "scanned")
    check("says OCR is required", "OCR" in res["note"], True)

    print("\na file that is not a PDF at all")
    junk = tmp / "junk.pdf"
    junk.write_bytes(b"this is not a PDF")
    res = M.extract_from_pdf(str(junk))
    check("status", res["status"], "error")

    batch(tmp)

    print()
    if FAILED:
        print(f"{len(FAILED)} check(s) failed: " + ", ".join(FAILED))
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
