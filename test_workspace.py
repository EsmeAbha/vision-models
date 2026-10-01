"""Isolated integration checks for the local workspace (no private fixtures)."""
import hashlib
import io
import json
import shutil
import uuid
import unittest
from unittest.mock import patch
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import workspace_server as server


def synthetic_pdf():
    lines = []
    for i, (label, a, b) in enumerate([
        ("Operating statement", "", ""),
        ("Rental income", "1200.00", "1300.00"),
        ("Other income", "200.00", "300.00"),
        ("Parking income", "100.00", "100.00"),
        ("Total income", "1500.00", "1700.00"),
    ]):
        y = 740 - i * 24
        for x, text in ((50, label), (300, a), (420, b)):
            if text:
                if x != 50:
                    x -= sum(3.336 if c == "." else 6.672 for c in text)
                lines.append(f"BT /F1 12 Tf {x} {y} Td ({text}) Tj ET")
    stream = "\n".join(lines).encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    pdf = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objs, 1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    start = len(pdf)
    pdf += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{n:010} 00000 n \n".encode() for n in offsets[1:])
    pdf += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()
    return pdf


class WorkspaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = server.ROOT / "outputs" / ("finai-test-" + uuid.uuid4().hex)
        server.DATA = cls.tmp / "outputs"
        server.INTAKE = cls.tmp / "intake"
        server.DATA.mkdir(parents=True)
        server.INTAKE.mkdir(parents=True)
        server.DB = server.DATA / "test.db"
        with server.connection() as db:
            db.execute("CREATE TABLE records (kind TEXT, id TEXT, body TEXT, PRIMARY KEY(kind,id))")
        cls.client = TestClient(server.app)
        cls.client.get("/")

    @classmethod
    def tearDownClass(cls):
        server.POOL.shutdown(wait=True)
        assert cls.tmp.resolve().is_relative_to((server.ROOT / "outputs").resolve())
        shutil.rmtree(cls.tmp)

    def upload(self, name, data):
        response = self.client.post("/api/upload", files=[("files", (name, data))])
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()[0]["id"]

    def plan(self, files, mode="extract"):
        r = self.client.post("/api/plan", json={"files": files, "prompt": "Extract financial tables and check totals", "mode": mode, "skill": "workbook-qc" if mode == "qc" else "table-extraction"})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def test_session_origin_and_path_boundaries(self):
        anonymous = TestClient(server.app)
        self.assertEqual(anonymous.get("/api/state").status_code, 401)
        self.assertEqual(self.client.post("/api/intake", json={"path": "../../"}).status_code, 400)
        self.assertEqual(self.client.post("/api/intake", json={"path": str(Path.cwd())}).status_code, 400)
        self.assertEqual(self.client.get("/api/state", headers={"Origin": "https://untrusted.example"}).status_code, 403)
        self.assertEqual(self.client.get("/api/state", headers={"Host": "untrusted.example"}).status_code, 403)
        self.assertEqual(self.client.post("/api/upload", files=[("files", ("worker.exe", b"bad"))]).status_code, 400)

    def test_real_pdf_to_workbook_and_hash(self):
        source = self.upload("synthetic-statement.pdf", synthetic_pdf())
        job = self.plan([source])
        self.assertEqual(job["status"], "Planned")
        server.process_job(job["id"])
        result = self.client.get("/api/jobs/" + job["id"]).json()
        self.assertIn(result["status"], ("Completed", "Needs review"), result)
        workbooks = [a for a in result["artifacts"] if a["kind"] == "Workbook"]
        self.assertEqual(len(workbooks), 1, result)
        output = self.client.get(f"/api/jobs/{job['id']}/download/{workbooks[0]['id']}")
        self.assertEqual(hashlib.sha256(output.content).hexdigest(), workbooks[0]["sha256"])
        wb = load_workbook(io.BytesIO(output.content))
        self.assertGreater(wb.active.max_row, 3)
        self.assertIn("Source page", [c.value for c in wb.active[1]])
        wb.close()
        listed = self.client.get("/api/state").json()["jobs"]
        self.assertNotIn("synthetic-statement.pdf", json.dumps(listed))
        self.assertEqual(self.client.get(f"/api/jobs/{job['id']}/download/unlisted.xlsx").status_code, 404)

    def test_qc_finds_cell_errors_and_does_not_certify_unmapped_source(self):
        wb = Workbook()
        wb.active["A1"] = "#REF!"
        stream = io.BytesIO()
        wb.save(stream)
        source = self.upload("source.pdf", synthetic_pdf())
        book = self.upload("prepared.xlsx", stream.getvalue())
        job = self.plan([source, book], "qc")
        server.process_job(job["id"])
        result = server.get("job", job["id"])
        self.assertEqual(result["status"], "Needs review")
        self.assertTrue(any(f["actual"] == "#REF!" for f in result["findings"]))
        self.assertTrue(any("mapping" in f["message"] for f in result["findings"]))
        self.assertFalse(any(a["kind"] == "Workbook" for a in result["artifacts"]))
        f = result["findings"][0]
        response = self.client.post(f"/api/jobs/{job['id']}/findings/{f['id']}", json={"action": "Acknowledge", "note": "Reviewed; correction required"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "Needs review")

    def test_plan_validation_and_safe_error(self):
        source = self.upload("broken.pdf", b"not a PDF")
        wrong = self.client.post("/api/plan", json={"files": [source], "prompt": "Review workbook", "mode": "qc", "skill": "workbook-qc"})
        self.assertEqual(wrong.status_code, 400)
        job = self.plan([source])
        server.process_job(job["id"])
        result = server.get("job", job["id"])
        self.assertEqual(result["status"], "Failed")
        self.assertNotIn(str(server.INTAKE), result["error"])
        self.assertNotIn("Traceback", result["error"])

    def test_draft_release_gates_and_workflow_compatibility(self):
        r = self.client.post("/api/studio/drafts", json={"name": "Synthetic workflow", "category": "workflow", "config": {"steps": ["export", "intake"]}})
        key = r.json()["id"]
        self.assertEqual(self.client.post(f"/api/studio/drafts/{key}/review").status_code, 400)
        result = self.client.post(f"/api/studio/drafts/{key}/test").json()
        self.assertFalse(result["test"]["passed"])
        r = self.client.post("/api/studio/drafts", json={"name": "Compatible workflow", "category": "workflow", "config": {"steps": ["intake", "parse", "extract", "validate", "map", "qc", "export"]}})
        key = r.json()["id"]
        self.assertTrue(self.client.post(f"/api/studio/drafts/{key}/test").json()["test"]["passed"])
        self.assertEqual(self.client.post(f"/api/studio/drafts/{key}/review").json()["status"], "In review")

    def test_utility_skill_runs_from_default_plan_and_publishes_preview(self):
        source = self.upload("bill.pdf", synthetic_pdf())
        response = self.client.post("/api/plan", json={"files": [source], "skill": "utility-bill-extraction"})
        self.assertEqual(response.status_code, 200, response.text)
        plan = response.json()
        self.assertIn("utility bill", plan["prompt"])
        self.assertIn("Account number", plan["output_fields"])
        sample = {"columns": ["Account number"], "rows": [{"values": ["001234"]}],
                  "evidence": [{"field": "Account number", "value": "001234", "destination": "Extract!A2", "matches": [{"page": 1, "region": "Text line 1", "quote": "Account 001234", "method": "Native document text"}]}],
                  "issues": [{"message": "Review the bill before use", "severity": "Warning"}],
                  "checks": [{"check": "Source match", "ok": True}],
                  "pages": [{"page": 1, "text": "Account 001234", "method": "Native document text"}]}
        with patch.object(server.field_skills, "extract", return_value=sample) as runner:
            server.process_job(plan["id"])
        runner.assert_called_once()
        result = server.get("job", plan["id"])
        self.assertEqual(result["status"], "Needs review")
        self.assertEqual(result["extractions"][0]["rows"][0]["values"], ["001234"])
        self.assertEqual(len(result["artifacts"]), 6)
        self.assertEqual(result["files"][0]["status"], "Needs review")
        image = self.upload("bill.png", b"synthetic-image")
        self.assertEqual(self.client.post("/api/plan", json={"files": [image], "skill": "utility-bill-extraction"}).status_code, 200)
        self.assertEqual(self.client.post("/api/plan", json={"files": [image], "skill": "table-extraction", "prompt": "Extract tables"}).status_code, 400)
        self.assertEqual(self.client.post("/api/plan", json={"files": [image], "skill": "extract-requested-fields"}).status_code, 400)


if __name__ == "__main__":
    unittest.main(verbosity=2)
