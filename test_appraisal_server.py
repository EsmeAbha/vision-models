"""Drive the running appraisal page over HTTP, the way the browser does.

Checks the parts the unit tests cannot: that an upload is accepted, that the
progress stream reports every file, that the preview endpoint returns the
table, and that the workbook comes back down the wire as a real workbook.

Start the server first, then:

    .\\.venv\\Scripts\\python.exe appraisal_server.py
    .\\.venv\\Scripts\\python.exe test_appraisal_server.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

from openpyxl import load_workbook

import test_mhc_rating_table as F

BASE = "http://127.0.0.1:7885"
FAILED = []

# Set when the server under test is password-protected, which it is whenever
# it was started by serve_appraisals.ps1 rather than run bare on loopback.
COOKIE = ""


def check(label, got, want):
    ok = got == want
    print(f"  {'PASS' if ok else 'FAIL'}  {label}")
    if not ok:
        print(f"        wanted {want!r}")
        print(f"        got    {got!r}")
        FAILED.append(label)


def get(path, raw=False):
    req = urllib.request.Request(BASE + path)
    if COOKIE:
        req.add_header("Cookie", COOKIE)
    with urllib.request.urlopen(req, timeout=120) as r:
        body = r.read()
    return body if raw else json.loads(body)


def sign_in():
    """Pick up a session if this server wants a password. -> message or None."""
    global COOKIE
    with urllib.request.urlopen(BASE + "/api/session", timeout=20) as r:
        state = json.loads(r.read())
    if not state.get("required"):
        return None

    password = os.environ.get("APPRAISAL_PASSWORD", "")
    if not password:
        return ("that server wants a password. Either run a bare one on "
                "loopback,\n  or set APPRAISAL_PASSWORD to the one it was "
                "started with:\n"
                '    $env:APPRAISAL_PASSWORD = "..."; '
                ".\\.venv\\Scripts\\python.exe test_appraisal_server.py")

    req = urllib.request.Request(
        BASE + "/api/login", data=json.dumps({"password": password}).encode(),
        headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=20) as r:
            COOKIE = (r.headers.get("set-cookie") or "").split(";")[0]
    except urllib.error.HTTPError as e:
        return f"the password was refused ({e.code})"
    print(f"signed in to the guarded server on {BASE}\n")
    return None


def post_zip(path, name="appraisals.zip"):
    """A multipart upload, hand-built so the test needs no extra library."""
    boundary = "----mhcboundary9173"
    body = io.BytesIO()
    body.write(f"--{boundary}\r\n".encode())
    body.write(f'Content-Disposition: form-data; name="file"; '
               f'filename="{name}"\r\n'.encode())
    body.write(b"Content-Type: application/zip\r\n\r\n")
    body.write(Path(path).read_bytes())
    body.write(f"\r\n--{boundary}--\r\n".encode())
    req = urllib.request.Request(
        BASE + "/api/jobs", data=body.getvalue(),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    if COOKIE:
        req.add_header("Cookie", COOKIE)
    with urllib.request.urlopen(req, timeout=300) as r:
        return json.loads(r.read())


def follow(jid):
    """Read the SSE stream to the end, returning the events in order."""
    events = []
    req = urllib.request.Request(BASE + f"/api/jobs/{jid}/events")
    if COOKIE:
        req.add_header("Cookie", COOKIE)
    with urllib.request.urlopen(req, timeout=600) as r:
        for line in r:
            line = line.decode("utf-8").strip()
            if not line.startswith("data: "):
                continue
            m = json.loads(line[6:])
            events.append(m)
            if m.get("type") in ("done", "error"):
                break
    return events


def build_archive(tmp):
    src = tmp / "archive"
    (src / "2026").mkdir(parents=True, exist_ok=True)
    F.build_with_guide(src / "sunset_ridge.pdf")
    F.build_with_guide(src / "2026" / "oak_meadows.pdf")
    F.build_without_guide(src / "cedar_flats.pdf")
    F.build_scan(src / "scanned_copy.pdf")
    archive = tmp / "appraisals.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        for p in src.rglob("*"):
            if p.is_file():
                zf.write(p, p.relative_to(src))
    return archive


def main():
    try:
        problem = sign_in()
        if problem:
            print(problem)
            return 2
        spec = get("/api/spec")
    except Exception as e:
        print(f"the server is not answering on {BASE} ({e})")
        print("start it:  .\\.venv\\Scripts\\python.exe appraisal_server.py")
        return 2

    print("the page describes what it looks for")
    check("four class columns", spec["columns"], F.M.CLASS_COLUMNS)
    check("ten core rows", len(spec["rows"]), 10)
    check("two star rows", len(spec["star_rows"]), 2)

    print("\nthe page itself is served")
    html = get("/", raw=True).decode("utf-8")
    check("the shell is there", "<div id=\"app\"" in html, True)
    check("it loads its script", "/assets/app.js" in html, True)
    check("the stylesheet is served",
          b"--accent" in get("/assets/style.css", raw=True), True)

    tmp = Path(tempfile.mkdtemp(prefix="mhc_http_"))
    archive = build_archive(tmp)

    print("\nuploading an archive of four appraisals")
    job = post_zip(archive)
    check("the job was accepted", bool(job.get("id")), True)

    events = follow(job["id"])
    kinds = [e["type"] for e in events]
    check("the stream starts then finishes", (kinds[0], kinds[-1]),
          ("start", "done"))
    files = [e for e in events if e["type"] == "file"]
    check("every file is reported once", len(files), 4)

    done = events[-1]
    check("two carried the guide", done["counts"]["found"], 2)
    check("one had none", done["counts"]["not_found"], 1)
    check("one was a scan", done["counts"]["scanned"], 1)
    check("a workbook was named", bool(done.get("workbook")), True)

    print("\nthe finished job can be read back")
    state = get(f"/api/jobs/{job['id']}")
    check("status", state["status"], "done")
    check("it lists all four files", len(state["files"]), 4)

    print("\nthe preview endpoint returns one table")
    idx = next(i for i, f in enumerate(state["files"])
               if f["status"] == "found")
    table = get(f"/api/jobs/{job['id']}/table/{idx}")
    check("the four class columns", table["columns"], F.M.CLASS_COLUMNS)
    rows = {r["category"]: r for r in table["rows"] if not r["section"]}
    check("all twelve rows", len(rows), 12)
    check("Density / Class A",
          rows["Density"]["values"][0], "Low (4-7 sites/acre)")
    check("Star Rating (Woodall) / Unratable",
          rows["Star Rating (Woodall)"]["values"][3], "N/A")
    check("the heading came through", table["title"], F.TITLE)

    print("\nasking for a table from a file that had no guide")
    miss = next(i for i, f in enumerate(state["files"])
                if f["status"] == "not_found")
    try:
        get(f"/api/jobs/{job['id']}/table/{miss}")
        check("it refuses", "no error", "a 404")
    except urllib.error.HTTPError as e:
        check("it refuses with 404", e.code, 404)

    print("\nthe workbook downloads and opens")
    blob = get(f"/api/jobs/{job['id']}/workbook", raw=True)
    out = tmp / "downloaded.xlsx"
    out.write_bytes(blob)
    wb = load_workbook(out)
    check("Index and Combined lead it", wb.sheetnames[:2],
          ["Index", "Combined"])
    check("a sheet per appraisal with the guide", len(wb.sheetnames), 4)
    body = [[c.value for c in r]
            for r in wb["Combined"].iter_rows(min_row=2)]
    check("24 rows across the two appraisals", len(body), 24)
    check("the index covers every file",
          len(list(wb["Index"].iter_rows(min_row=2))), 4)

    print("\nuploading something that is not a zip")
    notzip = tmp / "notes.txt"
    notzip.write_text("hello")
    try:
        post_zip(notzip, name="notes.txt")
        check("it refuses", "accepted", "rejected")
    except urllib.error.HTTPError as e:
        check("it refuses with 400", e.code, 400)

    print()
    if FAILED:
        print(f"{len(FAILED)} check(s) failed: " + ", ".join(FAILED))
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
