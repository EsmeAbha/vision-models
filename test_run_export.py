import os
import tempfile
import unittest

from openpyxl import load_workbook

import run_export


class FieldsWorkbookTests(unittest.TestCase):
    def build(self, rows, fields=None):
        path = os.path.join(tempfile.mkdtemp(), "fields.xlsx")
        run_export.fields_workbook([{"path": "", "file": "bill.pdf",
                                     "rows": rows}], path, fields=fields)
        wb = load_workbook(path)
        return [list(r) for r in wb["Fields"].iter_rows(values_only=True)], wb

    def test_a_field_empty_on_every_page_gets_no_column(self):
        rows = [
            {"field": "account_number", "value": "EL-1", "page": 1},
            {"field": "total_gas_bill", "value": "", "page": 1},
            {"field": "account_number", "value": "EL-1", "page": 2},
            {"field": "total_gas_bill", "value": "", "page": 2},
        ]
        # The type lists a field that was never ticked, too.
        sheet, wb = self.build(rows, fields=["provider_name", "account_number",
                                             "total_gas_bill"])
        self.assertEqual(sheet[0], run_export.LEAD + ["account_number"])
        self.assertEqual(len(sheet), 3)
        # What was not found is still on record, just not as a blank column.
        self.assertEqual(wb["Evidence"].max_row - 1, 4)

    def test_a_field_found_on_one_page_keeps_its_column_and_blank_elsewhere(self):
        rows = [
            {"field": "taxes", "value": "", "page": 1},
            {"field": "taxes", "value": "6.10", "page": 2},
        ]
        sheet, _ = self.build(rows)
        self.assertEqual(sheet[0][-1], "taxes")
        self.assertIn(sheet[1][-1], (None, ""))
        self.assertEqual(sheet[2][-1], "6.10")

    def test_the_type_order_is_kept_among_the_columns_left(self):
        rows = [{"field": f, "value": "x", "page": 1}
                for f in ("taxes", "account_number", "billing_date")]
        sheet, _ = self.build(rows, fields=["account_number", "provider_name",
                                            "billing_date", "taxes"])
        self.assertEqual(sheet[0][len(run_export.LEAD):],
                         ["account_number", "billing_date", "taxes"])



class SplitPathTests(unittest.TestCase):
    def test_a_windows_path_is_kept_as_windows_writes_it(self):
        self.assertEqual(
            run_export.split_path(r"C:\Users\Esme Abha\Downloads\APR\bill.pdf"),
            (r"C:\Users\Esme Abha\Downloads\APR\bill.pdf", "APR", "bill.pdf"))

    def test_no_path_leaves_path_and_folder_empty(self):
        self.assertEqual(run_export.split_path("", "bill.pdf"), ("", "", "bill.pdf"))


if __name__ == "__main__":
    unittest.main()
